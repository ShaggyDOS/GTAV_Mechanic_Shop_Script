#!/usr/bin/env python3
"""Build/test and package a documented QA candidate, without touching game saves."""
import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import zipfile

parser=argparse.ArgumentParser()
parser.add_argument('--label', required=True)
parser.add_argument('--issues', nargs='+', required=True)
parser.add_argument('--base', default='development/v0.3.3.1')
parser.add_argument('--artifacts-only', action='store_true')
parser.add_argument('--include-diagnostics', action='store_true')
parser.add_argument('--final-review', action='store_true')
args=parser.parse_args()
repo=Path(__file__).resolve().parents[1]
release=repo/f'docs/releases/v0.3.3.1-{args.label}'
release.mkdir(parents=True,exist_ok=True)
guide=release/'DEVELOPER_GUIDE.md'
if not guide.exists():raise SystemExit('Write issue-specific developer guide before packaging.')
sources=['VisualStudio/scripts/GlassHeroesBusiness.2.cs','Install/scripts/GlassHeroesBusiness.2.cs']
hashes={p:hashlib.sha256((repo/p).read_bytes()).hexdigest() for p in sources}
assert len(set(hashes.values()))==1,'Install/Visual Studio source mismatch'
assert hashlib.sha256((repo/'GlassHeroesBusiness_v0.3.3.0_RC1.2.cs').read_bytes()).hexdigest()=='030a44cfe4afc199c7865cd721361056be54398cf735eb1dbfc53cc87dee6949'

def run(command, filename):
    result=subprocess.run(command,cwd=repo,text=True,stdout=subprocess.PIPE,stderr=subprocess.STDOUT)
    (release/filename).write_text(result.stdout)
    if result.returncode:
        print(result.stdout)
        raise SystemExit(result.returncode)
    return result.stdout

diagnostic_counts = None
previous_validation = None
if args.artifacts_only:
    previous_validation=json.loads((release/'validation.json').read_text())
    assert previous_validation['source_sha256']==hashes, 'Source changed; rerun all validation'
    build=(release/'build.txt').read_text()
    tests=(release/'behavior-tests.txt').read_text()
    if args.include_diagnostics:
        diagnostics=(release/'diagnostics-tests.txt').read_text()
else:
    with tempfile.TemporaryDirectory(prefix='glassheroes-build-') as tmp:
        build=run(['dotnet','build','VisualStudio/GlassHeroesBusiness.csproj','--configuration','Release',
            '--property:BaseIntermediateOutputPath='+tmp+'/obj/','--property:OutputPath='+tmp+'/bin/',
            '--maxcpucount:2','--verbosity','minimal'],'build.txt')
    checks=run([sys.executable,'verification/run_static_checks.py',str(release)],'static-checks.txt')
    tests=run([sys.executable,'verification/run_issue_tests.py',*args.issues],'behavior-tests.txt')
    if args.include_diagnostics:
        diagnostics=run([sys.executable,'verification/run_diagnostics_tests.py'],'diagnostics-tests.txt')
    if args.label.startswith('qa'):
        run([sys.executable,'verification/check_preservation.py',str(release/'preservation.json')],'preservation-checks.txt')
static=json.loads((release/'StaticVerification.json').read_text())
counts=re.findall(r'#(\d+) (\d+)/(\d+) PASS;',tests)
assert set(x[0] for x in counts)==set(args.issues), 'Missing executed issue tests'
assert all(int(p)==int(t) and int(t)>0 for _,p,t in counts),'Incomplete test execution'
if args.include_diagnostics:
    match=re.search(r'(\d+)/(\d+) diagnostic behavior tests passed',diagnostics)
    assert match and int(match[1])==int(match[2]) and int(match[2])>0, 'Diagnostics did not execute fully'
    diagnostic_counts={'passed':int(match[1]),'total':int(match[2])}
patch=subprocess.check_output(['git','diff',args.base,'--',*sources],cwd=repo)
(release/'source.diff').write_bytes(patch)
with tempfile.TemporaryDirectory(prefix='glassheroes-patch-') as tmp:
    root=Path(tmp)
    for source in sources:
        p=root/source;p.parent.mkdir(parents=True,exist_ok=True)
        p.write_bytes(subprocess.check_output(['git','show',args.base+':'+source],cwd=repo))
    subprocess.run(['git','apply','--check',str(release/'source.diff')],cwd=root,check=True)
    subprocess.run(['git','apply',str(release/'source.diff')],cwd=root,check=True)
    assert all((root/p).read_bytes()==(repo/p).read_bytes() for p in sources)
run(['git','diff','--check',args.base],'whitespace-check.txt')
validation={
 'version':'v0.3.3.1 EXPERIMENTAL / QA ONLY','label':args.label,'included_issues':(['1'] if args.include_diagnostics else [])+args.issues,
 'created_utc':previous_validation['created_utc'] if previous_validation else datetime.now(timezone.utc).isoformat(),
 'base_commit':subprocess.check_output(['git','rev-parse',args.base],cwd=repo,text=True).strip(),
 'sdk':subprocess.check_output(['dotnet','--version'],text=True).strip(),
 'target':'net48 / C# 7.3 / ScriptHookVDotNet2 2.11.6',
 'release_build':'PASS','production_warnings':sorted(set(re.findall(r'warning (CS[0-9]+)',build))),
 'static_checks':{'passed':static['passed'],'total':static['total']},
 'behavior_tests':{issue:{'passed':int(p),'total':int(t)} for issue,p,t in counts},
 'source_sha256':hashes,'source_copies_equal':True,'source_diff_applies_exactly':True,
 'historical_rc1_unchanged':True,'business_saves_or_live_settings_touched':False,
 'gta_testing':'NOT RUN','windows_build':'NOT RUN','crash_cause':'UNPROVEN',
 'review_status':'PENDING; no integration or main merge authorized by this artifact',
}
if diagnostic_counts:
    validation['behavior_tests']['1']=diagnostic_counts
    validation['pr9_diagnostics_commit']=json.loads((repo/'verification/pr9_diagnostics_reference.json').read_text())['commit']
validation['behavior_test_total']=sum(v['total'] for v in validation['behavior_tests'].values())
validation['behavior_test_passed']=sum(v['passed'] for v in validation['behavior_tests'].values())
validation['final_review_only']=args.final_review
(release/'validation.json').write_text(json.dumps(validation,indent=2)+'\n')
text=guide.read_text().split('\n## Recorded automated results\n')[0]
text+='\n## Recorded automated results\n\n'
text+=f'Real Release build: PASS. SDK {validation["sdk"]}; net48 / C# 7.3 / SHVDN2 2.11.6. Two existing production warnings: CS0649 and CS0169.\n\n'
text+=f'Existing static checks: {static["passed"]}/{static["total"]} PASS. '
text+='; '.join(f'Issue #{i}: {p}/{t} behavioral tests PASS' for i,p,t in counts)+'.\n\n'
if diagnostic_counts:
    text+=f'PR #9 diagnostic logger tests: {diagnostic_counts["passed"]}/{diagnostic_counts["total"]} PASS; total behavioral tests: {validation["behavior_test_passed"]}/{validation["behavior_test_total"]} PASS.\n\n'
text+='Install and Visual Studio sources have equal SHA256: '+next(iter(hashes.values()))+'.\n\n'
text+='Full source.diff applies to the recorded integration baseline and reproduces both sources byte-for-byte. Build, static and behavior transcripts plus validation.json are included. The static runner compiled=false field describes its static-only run; actual compilation is separately verified. Test harnesses use extracted production members with GTA APIs stubbed, not real game sessions. GTA testing: NOT RUN. Windows build: NOT RUN. Stack-overflow cause: UNPROVEN. Review: PENDING.\n'
guide.write_text(text)

from docx import Document
from docx.shared import Pt
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Preformatted
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from xml.sax.saxutils import escape
from pypdf import PdfReader
import textwrap
stem=f'GlassHeroesBusiness_v0.3.3.1_{args.label}_Developer_Guide'
doc=Document();doc.styles['Normal'].font.size=Pt(10)
styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='CodeSmall',fontName='Courier',fontSize=7,leading=9))
flow=[];code=False
for line in text.splitlines():
    if line.startswith('```'):code=not code;continue
    if code:
        p=doc.add_paragraph();r=p.add_run(line);r.font.name='Courier New';r.font.size=Pt(8)
        for part in textwrap.wrap(line,100,replace_whitespace=False,drop_whitespace=False) or ['']:
            flow.append(Preformatted(part,styles['CodeSmall']))
    elif not line:flow.append(Spacer(1,5))
    elif line.startswith('#'):
        n=len(line)-len(line.lstrip('#'));heading=line[n:].strip()
        doc.add_heading(heading,level=min(n-1,3));flow.append(Paragraph(escape(heading),styles['Title' if n==1 else 'Heading2']))
    else:doc.add_paragraph(line);flow.append(Paragraph(escape(line),styles['BodyText']))
doc.save(release/(stem+'.docx'))
def footer(c,d):
    c.setFont('Helvetica',8);c.drawString(42,22,'EXPERIMENTAL / REVIEW PENDING | GTA testing NOT RUN');c.drawRightString(570,22,str(d.page))
SimpleDocTemplate(str(release/(stem+'.pdf')),pagesize=(612,792),leftMargin=42,rightMargin=42,topMargin=40,bottomMargin=38).build(flow,onFirstPage=footer,onLaterPages=footer)
pdf='\n'.join(p.extract_text() for p in PdfReader(release/(stem+'.pdf')).pages)
docx='\n'.join(p.text for p in Document(release/(stem+'.docx')).paragraphs)
for phrase in ['EXPERIMENTAL','NOT RUN','UNPROVEN','Save.ini','Recorded automated results']:
    assert phrase in pdf and phrase in docx,('Missing guide qualification',phrase)

package_dir=repo/'build/releases';package_dir.mkdir(parents=True,exist_ok=True)
package=package_dir/f'GlassHeroesBusiness_v0.3.3.1_{args.label}_EXPERIMENTAL.zip'
files=[repo/p for p in [*sources,'VisualStudio/GlassHeroesBusiness.csproj','VisualStudio/GlassHeroesBusiness.sln','VisualStudio/START_HERE.md','README.md','RELEASE_NOTES.md',
 'Install/scripts/GlassHeroesBusiness/Settings.ini.example','AGENTS.md','docs/CODEX_V0_3_3_1_HANDOFF.md',
 'docs/V0_3_3_1_TEST_PLAN.md','GlassHeroesBusiness_v0.3.3.0_RC1.2.cs','BASELINE_SHA256.txt',
 'SourceDiff_v0.3.2.9a_to_v0.3.3.0_RC1.patch','Verification/static_verify.py',
 'verification/run_issue_tests.py','verification/run_static_checks.py','verification/prepare_experimental_package.py']]
files += [repo/f'verification/issue{i}_tests.cs' for i in args.issues]
if args.label.startswith('qa'):
    files += [repo/'verification/check_preservation.py']
    if not args.final_review:
        for issue in args.issues:
            files += sorted(p for p in (repo/f'docs/releases/v0.3.3.1-issue{issue}').iterdir() if p.is_file())
if args.include_diagnostics:
    files += [repo/'verification/run_diagnostics_tests.py',repo/'verification/diagnostics_tests.cs',repo/'verification/pr9_diagnostics_reference.json']
files += sorted(p for p in release.rglob('*') if p.is_file())
with zipfile.ZipFile(package,'w',zipfile.ZIP_DEFLATED) as z:
    for file in files:
        arcname=str(file.relative_to(repo))
        if args.final_review and arcname=='GlassHeroesBusiness_v0.3.3.0_RC1.2.cs':
            arcname='verification/baselines/GlassHeroesBusiness_v0.3.3.0_RC1.2.cs.txt'
        z.write(file,arcname)
    if args.label.startswith('qa'):
        z.write(release/'READ_ME_FIRST.txt','READ_ME_FIRST.txt')
with zipfile.ZipFile(package) as z:
    assert z.testzip() is None
    assert len(z.namelist())==len(set(z.namelist())), 'Duplicate archive paths'
    if args.final_review:
        assert set(n for n in z.namelist() if n.endswith('.2.cs'))==set(sources), 'Unexpected active business script'
        assert not any(n.endswith(('/Save.ini','/Settings.ini','.dll','.zip')) for n in z.namelist()), 'Unexpected live configuration, binary or nested archive'
    assert all(z.read(p)==(repo/p).read_bytes() for p in sources)
manifest={'package':package.name,'sha256':hashlib.sha256(package.read_bytes()).hexdigest(),
 'zip_crc':'PASS','sources_match':'PASS','guide_formats':'MD, DOCX, PDF verified',
 'included_issues':validation['included_issues'],'behavior_test_total':validation['behavior_test_total'],'gta_testing':'NOT RUN','review_status':'PENDING'}
(package.with_suffix('.manifest.json')).write_text(json.dumps(manifest,indent=2)+'\n')
print(json.dumps(validation,indent=2));print('Verified package:',package.name)
