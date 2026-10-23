#!/usr/bin/env python3
"""Package already executed Phase 1 results; never installs, changes saves or merges."""
from pathlib import Path
import difflib,hashlib,json,re,subprocess,tempfile,zipfile
from datetime import datetime,timezone
from xml.sax.saxutils import escape
import textwrap
from docx import Document
from docx.shared import Pt
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Preformatted
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from pypdf import PdfReader
repo=Path(__file__).resolve().parents[2]
release=repo/'docs/releases/v0.3.3.1-shvdn3-phase1'
source='experimental/shvdn3/Install/scripts/GlassHeroesBusiness.3.cs'
project='experimental/shvdn3/VisualStudio/GlassHeroesBusiness.V3.Experimental.csproj'
solution='experimental/shvdn3/VisualStudio/GlassHeroesBusiness.V3.Experimental.sln'
raw=(repo/source).read_bytes();sha=hashlib.sha256(raw).hexdigest()
parity=json.loads((release/'source-parity.json').read_text());assert parity['v3_sha256']==sha
static2=json.loads((release/'StaticVerification.json').read_text());static3=json.loads((release/'v3-static/StaticVerification.json').read_text())
assert static3['source_sha256']==sha
assert static2['passed']==static2['total']==static3['passed']==static3['total']==21
for name in ['baseline-v2-build.txt','v3-build.txt']:
 s=(release/name).read_text();assert 'Build succeeded.' in s and '0 Error(s)' in s,name
counts={}
for generation in ['baseline-v2','v3']:
 corrective=re.findall(r'#(\d+) (\d+)/(\d+) PASS;', (release/(generation+'-behavior.txt')).read_text())
 assert set(x[0] for x in corrective)==set('234567')
 assert all(int(p)==int(t) for _,p,t in corrective)
 count=sum(int(t) for _,_,t in corrective);assert count==54
 logger=re.search(r'(\d+)/(\d+) diagnostic behavior tests passed',(release/(generation+'-logger.txt')).read_text());assert logger and logger[1]==logger[2]=='9'
 counts[generation]={'corrective':54,'logger':9,'inherited_total':63}
migration=re.search(r'Migration (\d+)/(\d+) PASS;', (release/'migration-behavior.txt').read_text());assert migration and migration[1]==migration[2]=='22'
api=json.loads((release/'api-metadata.json').read_text());deps=json.loads((release/'dependencies.json').read_text())
assert api['control_value_changes']==[] and api['v2_control_names_checked']==363 and api['native_hashes_checked']==43
assert not any(x['name']=='ScriptHookVDotNet2' for x in api['built_product']['references'])
assert deps['version']=='3.6.0'
base='7693ed2167405f475dd4d36392c5df1f257fe178'
full=subprocess.check_output(['git','diff',base,'--',source,project,solution],cwd=repo)
(release/'source.diff').write_bytes(full)
with tempfile.TemporaryDirectory(prefix='gh-migration-diff-') as temp:
 subprocess.run(['git','apply','--check',str(release/'source.diff')],cwd=temp,check=True)
 subprocess.run(['git','apply',str(release/'source.diff')],cwd=temp,check=True)
 for p in [source,project,solution]:assert (Path(temp)/p).read_bytes()==(repo/p).read_bytes(),p
fixture=(repo/'verification/migration/qa-r2-reference.cs.txt').read_text(encoding='utf-8-sig')
comparison=''.join(difflib.unified_diff(fixture.splitlines(True),raw.decode('utf-8-sig').splitlines(True),fromfile='QA_R2/GlassHeroesBusiness.2.cs.txt',tofile=source))
(release/'api-source.diff').write_text(comparison)
validation={'version':'v0.3.3.1 SHVDN3 Phase 1 EXPERIMENTAL','issue':20,'draft_pr':21,'branch':'migration/qa-r2-shvdn3-prep','r2_baseline_commit':base,'created_utc':datetime.now(timezone.utc).isoformat(),'sdk':subprocess.check_output(['dotnet','--version'],text=True).strip(),'target':'net48 / C# 7.3','reference':'ScriptHookVDotNet3 3.6.0','v2_release_build':'PASS (0 errors / 2 existing warnings)','v3_release_build':'PASS (0 errors / 2 existing warnings)','inherited_behavior':counts,'additional_migration_cases':22,'v3_total_behavior':85,'static_per_generation':{'passed':21,'total':21},'source_sha256':sha,'project_sha256':hashlib.sha256((repo/project).read_bytes()).hexdigest(),'solution_sha256':hashlib.sha256((repo/solution).read_bytes()).hexdigest(),'one_installation_source_linked_by_project':True,'whole_source_r2_parity_after_reviewed_api_reversal':'PASS','additive_source_diff_applies_exactly':True,'v2_sources_project_solution_unchanged':True,'save_config_schema_changes':False,'live_game_saves_settings_touched':False,'pr18_pr19_pr22_features_integrated':False,'gta_testing':'NOT RUN','windows_execution':'NOT RUN','enhanced_3_9_0_6_runtime_compatibility':'NOT VERIFIED','crash_cause':'UNATTRIBUTED / UNPROVEN','review_status':'PENDING; DRAFT; no merge or deployment'}
receipt=release/'extracted-archive-checks.json'
if receipt.exists():
 extracted=json.loads(receipt.read_text())
 assert extracted['source_sha256']==sha and extracted['project_sha256']==validation['project_sha256']
 assert extracted['release_build'].startswith('PASS') and extracted['source_parity']=='PASS'
 validation['clean_extracted_package_rebuild']='PASS (0 errors / 2 existing warnings)'
 validation['clean_extracted_package_static']='21/21 PASS'
 validation['clean_extracted_package_parity']='PASS'
(release/'validation.json').write_text(json.dumps(validation,indent=2)+'\n')
guide=release/'DEVELOPER_GUIDE.md';text=guide.read_text().split('\n## Recorded automated results\n')[0]
text+='\n## Recorded automated results\n\n'
text+='Real Release builds: v2 PASS, v3 PASS; both 0 errors and two inherited warnings (CS0169, CS0649). SDK '+validation['sdk']+'; product net48 / C#7.3.\n\n'
text+='Baseline v2: 54 corrective + 9 logger = 63/63 behavior PASS; 21/21 static PASS. Migrated v3: same 63/63 inherited behavior PASS plus 22/22 migration cases = 85/85 behavior PASS; 21/21 static PASS. The nine logger cases are counted once per generation.\n\n'
text+='Whole-source R2 parity after the reviewed API reversal: PASS. Financial/clock/save bodies and key order unchanged. Original v2 sources/project/solution unchanged. Project compiles the single installation file directly. Additive source.diff applies and reproduces source/project/solution exactly. Offline metadata: 363 existing Control names/values match, all 43 distinct used native hashes match, font index matches; compiled product has no SHVDN2 reference.\n\n'
if receipt.exists(): text+='Clean ZIP extraction: Release rebuild PASS (0 errors / 2 existing warnings), 21/21 static PASS, source parity PASS; all archive file hashes verified. Source, project and solution match this candidate.\n\n'
text+='Installation/source SHA256: '+sha+'.\n\n'
text+='GTA gameplay testing: NOT RUN. Windows execution: NOT RUN. Actual Enhanced 3.9.0.6 runtime binding/engine/save behavior: NOT VERIFIED. Crash cause: UNATTRIBUTED / UNPROVEN. These automated results authorize no merge, installation or stability claim.\n'
guide.write_text(text)
stem='GlassHeroesBusiness_v0.3.3.1_SHVDN3_Phase1_Developer_Report'
doc=Document();doc.styles['Normal'].font.size=Pt(10)
styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='SmallCode',fontName='Courier',fontSize=7,leading=9));flow=[];code=False
for line in text.splitlines():
 if line.startswith('```'):code=not code;continue
 if code:
  p=doc.add_paragraph();run=p.add_run(line);run.font.name='Courier New';run.font.size=Pt(8)
  for part in textwrap.wrap(line,96,replace_whitespace=False,drop_whitespace=False) or ['']:flow.append(Preformatted(part,styles['SmallCode']))
 elif not line:flow.append(Spacer(1,5))
 elif line.startswith('#'):
  n=len(line)-len(line.lstrip('#'));label=line[n:].strip();doc.add_heading(label,level=min(n-1,3));flow.append(Paragraph(escape(label),styles['Title' if n==1 else 'Heading2']))
 else:doc.add_paragraph(line);flow.append(Paragraph(escape(line),styles['BodyText']))
doc.save(release/(stem+'.docx'))
def footer(canvas,page):
 canvas.setFont('Helvetica',8);canvas.drawString(42,22,'SHVDN3 PHASE 1 EXPERIMENTAL | GTA testing NOT RUN | REVIEW PENDING');canvas.drawRightString(570,22,str(page.page))
SimpleDocTemplate(str(release/(stem+'.pdf')),pagesize=(612,792),leftMargin=42,rightMargin=42,topMargin=40,bottomMargin=38).build(flow,onFirstPage=footer,onLaterPages=footer)
pdf='\n'.join(p.extract_text() for p in PdfReader(release/(stem+'.pdf')).pages)
word='\n'.join(p.text for p in Document(release/(stem+'.docx')).paragraphs)
for phrase in ['EXPERIMENTAL','NOT RUN','UNATTRIBUTED','Save.ini','Recorded automated results','85/85']:
 assert phrase in pdf and phrase in word,phrase
out=repo/'build/releases';out.mkdir(exist_ok=True,parents=True)
archive=out/'GlassHeroesBusiness_v0.3.3.1_SHVDN3_Phase1_EXPERIMENTAL.zip'
files={repo/p for p in [source,project,solution,'experimental/shvdn3/README.md','experimental/shvdn3/VisualStudio/START_HERE.md','AGENTS.md','docs/coordination/CODEX_QA_R2_SHVDN3_MIGRATION.md','Verification/static_verify.py','SourceDiff_v0.3.2.9a_to_v0.3.3.0_RC1.patch','BASELINE_SHA256.txt','verification/run_issue_tests.py','verification/run_diagnostics_tests.py','verification/run_static_checks.py','verification/diagnostics_tests.cs','verification/pr9_diagnostics_reference.json']}
files.update(repo/f'verification/issue{i}_tests.cs' for i in range(2,8))
files.update(p for p in (repo/'verification/migration').rglob('*') if p.is_file() and not any(x in {'bin','obj','__pycache__'} for x in p.parts))
files.update(p for p in release.rglob('*') if p.is_file() and p.name!='FILES_SHA256.json')
manifest={'source_sha256':sha,'files':{str(p.relative_to(repo)):hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(files)}}
(release/'FILES_SHA256.json').write_text(json.dumps(manifest,indent=2)+'\n');files.add(release/'FILES_SHA256.json')
with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
 for p in sorted(files):z.write(p,str(p.relative_to(repo)))
 z.write(repo/'GlassHeroesBusiness_v0.3.3.0_RC1.2.cs','verification/baselines/GlassHeroesBusiness_v0.3.3.0_RC1.2.cs.txt')
 z.write(release/'READ_ME_FIRST.txt','READ_ME_FIRST.txt')
with zipfile.ZipFile(archive) as z:
 assert z.testzip() is None
 assert len(z.namelist())==len(set(z.namelist()))
 assert [p for p in z.namelist() if p.endswith('.3.cs')]==[source]
 assert not any(p.endswith(('.2.cs','.dll','.zip','/Save.ini','/Settings.ini')) for p in z.namelist())
 assert z.read(source)==raw
manifest={'package':archive.name,'sha256':hashlib.sha256(archive.read_bytes()).hexdigest(),'one_business_3_cs':True,'no_active_v2_or_live_save_settings_or_dll':True,'v3_behavior':'85/85 PASS','static_per_generation':'21/21 PASS','v2_release_build':'PASS','v3_release_build':'PASS','gta_testing':'NOT RUN','runtime_compatibility':'NOT VERIFIED','review_status':'PENDING / DRAFT'}
archive.with_suffix('.manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
print(json.dumps(validation,indent=2));print('Verified review package:',archive)
