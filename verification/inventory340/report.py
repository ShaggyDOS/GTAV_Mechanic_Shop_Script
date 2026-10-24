#!/usr/bin/env python3
"""Generate versioned development Word/PDF reports from the reviewed Markdown."""
from pathlib import Path
from xml.sax.saxutils import escape
import json
from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Preformatted,Table,TableStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'docs/features/v0.3.4'
inputs=['DEVELOPER_GUIDE.md','CHANGELOG.md','GTA_TEST_CHECKLIST.md','phaseE/V035_INVENTORY_READINESS.md']
fonts=Path('/usr/share/fonts/truetype/dejavu')
for name,file in [('Report','DejaVuSans.ttf'),('ReportBold','DejaVuSans-Bold.ttf'),('ReportItalic','DejaVuSans-Oblique.ttf'),('ReportMono','DejaVuSansMono.ttf')]:pdfmetrics.registerFont(TTFont(name,str(fonts/file)))
doc=Document();doc.styles['Normal'].font.name='Calibri';doc.styles['Normal'].font.size=Pt(10)
styles=getSampleStyleSheet()
for level in range(1,4):styles['Heading'+str(level)].fontName='ReportBold'
styles.add(ParagraphStyle(name='ReportBody',fontName='Report',fontSize=9,leading=12,spaceAfter=6))
styles.add(ParagraphStyle(name='ReportCell',fontName='Report',fontSize=7.5,leading=10))
styles.add(ParagraphStyle(name='ReportCode',fontName='ReportMono',fontSize=7,leading=9,spaceAfter=6))
story=[];table_rows=[]
def plain(text):return text.replace('`','').replace('**','')
def flush_table():
 if not table_rows:return
 table=doc.add_table(rows=0,cols=len(table_rows[0]));table.style='Table Grid'
 for row in table_rows:
  cells=table.add_row().cells
  for i,cell in enumerate(row):cells[i].text=cell
 data=[[Paragraph(escape(cell),styles['ReportCell']) for cell in row] for row in table_rows]
 pdf_table=Table(data,colWidths=[(A4[0]-72)/len(data[0])]*len(data[0]),repeatRows=1,hAlign='LEFT')
 pdf_table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e7edf4')),('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),.3,colors.grey),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5)]))
 story.append(pdf_table);story.append(Spacer(1,8));table_rows.clear()
for filename in inputs:
 code=False;buffer=[]
 for line in (OUT/filename).read_text().splitlines():
  if not code and line.startswith('|'):
   cells=[plain(x.strip()) for x in line.strip('|').split('|')]
   if not all(set(x)<=set('-: ') for x in cells):table_rows.append(cells)
   continue
  flush_table()
  if line.startswith('```'):
   if code:
    body='\n'.join(buffer);p=doc.add_paragraph();run=p.add_run(body);run.font.name='Consolas';run.font.size=Pt(8);story.append(Preformatted(body,styles['ReportCode']));buffer=[]
   code=not code;continue
  if code:buffer.append(line);continue
  if not line.strip():story.append(Spacer(1,4));continue
  if line.startswith('#'):
   level=len(line)-len(line.lstrip('#'));body=line[level:].strip();doc.add_heading(body,min(level,4));story.append(Paragraph(escape(body),styles['Heading'+str(min(level,3))]));continue
  body=plain(line);doc.add_paragraph(body);story.append(Paragraph(escape(body),styles['ReportBody']))
 flush_table()
word=OUT/'GlassHeroesBusiness_v0.3.4_Development_Report.docx';pdf=OUT/'GlassHeroesBusiness_v0.3.4_Development_Report.pdf'
doc.core_properties.title='Glass Heroes v0.3.4 Development Review';doc.core_properties.subject='EXPERIMENTAL — GTA testing not run';doc.save(word)
def footer(canvas,document):
 canvas.saveState();canvas.setFont('Report',7);canvas.setFillColor(colors.grey);canvas.drawString(36,22,'Glass Heroes v0.3.4 development review — GTA NOT RUN');canvas.drawRightString(A4[0]-36,22,str(document.page));canvas.restoreState()
SimpleDocTemplate(str(pdf),pagesize=A4,rightMargin=36,leftMargin=36,topMargin=40,bottomMargin=36).build(story,onFirstPage=footer,onLaterPages=footer)
(OUT/'REPORT_GENERATION.json').write_text(json.dumps({'inputs':inputs,'outputs':[word.name,pdf.name],'mode':'development documentation only','pdf_fonts':'embedded DejaVu Sans','gta_testing':'NOT RUN'},indent=2)+'\n')
print(word);print(pdf)
