#!/usr/bin/env python3
"""Generate Issue37 review documents and prove the source diff reproduces code."""
import hashlib,json,subprocess,sys,tempfile,textwrap
from pathlib import Path
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Pt
from pypdf import PdfReader
from reportlab.lib.styles import ParagraphStyle,getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph,Preformatted,SimpleDocTemplate,Spacer
ROOT=Path(__file__).resolve().parents[2];sys.path.insert(0,str(ROOT/'verification/maintenance36'))
from source_tools import SOURCE,members,strip_literals
BASE='75e3623f608a05d1a2b59d12c243dbcdfc061358';CODE='f6b90560304033ca538b585e61f850ee50863e10'
OUT=ROOT/'docs/features/issue37-placement-parking'
v=json.loads((OUT/'final-validation/validation.json').read_text());sha=hashlib.sha256(SOURCE.read_bytes()).hexdigest()
assert sha==v['source_sha256'] and v['behavior']['total']==1057 and v['build']['errors']==0
assert SOURCE.read_bytes()==subprocess.check_output(['git','show',CODE+':'+str(SOURCE.relative_to(ROOT))],cwd=ROOT)
(OUT/'source.diff').write_bytes(subprocess.check_output(['git','diff',BASE,CODE,'--',str(SOURCE.relative_to(ROOT))],cwd=ROOT))
with tempfile.TemporaryDirectory(prefix='gh37-diff-') as directory:
    p=Path(directory)/SOURCE.relative_to(ROOT);p.parent.mkdir(parents=True)
    old=subprocess.check_output(['git','show',BASE+':'+str(SOURCE.relative_to(ROOT))],cwd=ROOT);p.write_bytes(old)
    subprocess.run(['git','apply',str(OUT/'source.diff')],cwd=directory,check=True);assert p.read_bytes()==SOURCE.read_bytes()
before,after=members(old.decode()),members(SOURCE.read_text())
excerpts='# Issue37 before/after excerpts\n\nComplete exact patch: source.diff. Base and code checkpoint are identified in the developer guide.\n'
for name in ['UpdateVehicleParking','ConfigureRoutePoint','SetShopLocationFromPlayer']:
    excerpts+='\n## '+name+'\n'
    for label,body in [('Before',before[name]),('After',after[name])]:excerpts+='\n'+label+'\n\n```csharp\n'+body+'\n```\n'
for name in ['IssueDirectionFreeHandoff','PlacementCategoryVisible','TryCommitPlacementBytes']:
    excerpts+='\n## New helper: '+name+'\n\nBefore: no corresponding helper. After:\n\n```csharp\n'+after[name]+'\n```\n'
(OUT/'SOURCE_EXCERPTS.md').write_text(excerpts)
matrix='# Issue37 production navigation inventory\n\nSource-derived from the final actual navigation harness; GTA navigation/rendering NOT RUN.\n\n'
footers='# Issue37 production row descriptions\n\nSource-derived, not native visual/controller qualification.\n\n'
for line in (OUT/'final-validation/navigation.txt').read_text().splitlines():
    if line.startswith('MATRIX|'):
        _,page,title,parent,row,items=line.split('|',5)
        matrix+=f'## {page} — {title}\n\nHarness return: {parent}, row {row}.\n\n'+'\n'.join('- '+str(i)+': '+item for i,item in enumerate(items.split('; ')))+'\n\n'
    elif line.startswith('FOOTER|'):
        _,page,row,label,description=line.split('|',4);footers+=f'- {page} [{row}] {label}: {description}\n'
(OUT/'NAVIGATION_MATRIX.md').write_text(matrix.rstrip()+'\n');(OUT/'MENU_FOOTERS.md').write_text(footers)
(OUT/'SOURCE_SHA256.txt').write_text(sha+'  Install/scripts/GlassHeroesBusiness.3.cs\n')
documents=['DEVELOPER_GUIDE.md','INSTALL_AND_ROLLBACK.md','GTA_TEST_CHECKLIST.md','CHANGELOG.md','SOURCE_EXCERPTS.md']
content='\n\n'.join((OUT/name).read_text() for name in documents)
for name,file in [('Issue37Text','DejaVuSans.ttf'),('Issue37Bold','DejaVuSans-Bold.ttf'),('Issue37Mono','DejaVuSansMono.ttf')]:pdfmetrics.registerFont(TTFont(name,'/usr/share/fonts/truetype/dejavu/'+file))
styles=getSampleStyleSheet()
for name in ['Normal','BodyText','Title','Heading2']:styles[name].fontName='Issue37Bold' if name in ['Title','Heading2'] else 'Issue37Text'
styles.add(ParagraphStyle(name='Issue37Code',fontName='Issue37Mono',fontSize=7,leading=9))
doc=Document();doc.styles['Normal'].font.size=Pt(10);flow=[];code=False
for line in content.splitlines():
    if line.startswith('```'):code=not code;continue
    if code:
        run=doc.add_paragraph().add_run(line);run.font.name='Courier New';run.font.size=Pt(8)
        for part in textwrap.wrap(line,94,replace_whitespace=False,drop_whitespace=False) or ['']:flow.append(Preformatted(part,styles['Issue37Code']))
    elif not line:flow.append(Spacer(1,5))
    elif line.startswith('#'):
        n=len(line)-len(line.lstrip('#'));title=line[n:].strip();doc.add_heading(title,level=min(n-1,3));flow.append(Paragraph(escape(title),styles['Title' if n==1 else 'Heading2']))
    else:doc.add_paragraph(line);flow.append(Paragraph(escape(line),styles['BodyText']))
stem='GlassHeroes_v0.3.4_Issue37_Placement_Parking_Experimental_Developer_Guide';doc.save(OUT/(stem+'.docx'))
def footer(canvas,document):
    canvas.setFont('Issue37Text',8);canvas.drawString(42,22,'Issue37 EXPERIMENTAL | acceptance PENDING YES/NO | GTA NOT RUN');canvas.drawRightString(570,22,str(document.page))
pdf=OUT/(stem+'.pdf');SimpleDocTemplate(str(pdf),pagesize=(612,792),leftMargin=42,rightMargin=42,topMargin=40,bottomMargin=38).build(flow,onFirstPage=footer,onLaterPages=footer)
reader=PdfReader(pdf)
for text in ['\n'.join(p.text for p in Document(OUT/(stem+'.docx')).paragraphs),'\n'.join(p.extract_text() for p in reader.pages)]:
    for term in [sha,CODE,'1057/1057','Save.ini','Settings.ini','YES/NO','NOT RUN','UNATTRIBUTED']:assert term in text,term
r={'code_checkpoint':CODE,'base':BASE,'source_sha256':sha,'full_source_diff_reproduced':True,'word_pdf_text_verified':True,'pdf_pages':len(reader.pages),'gta_testing':'NOT RUN','owner_final_acceptance':'PENDING YES/NO','rc2_combination':False}
(OUT/'REPORT_RECEIPT.json').write_text(json.dumps(r,indent=2)+'\n');print(json.dumps(r,indent=2))
