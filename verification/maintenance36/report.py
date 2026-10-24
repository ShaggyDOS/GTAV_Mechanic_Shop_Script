#!/usr/bin/env python3
"""Generate Issue #36 review documents; never install, publish or accept a build."""
import hashlib
import json
import subprocess
import tempfile
import textwrap
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Pt
from pypdf import PdfReader
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, Preformatted, SimpleDocTemplate, Spacer

from source_tools import BASE, ROOT, SOURCE, baseline, members

OUT = ROOT / 'docs/features/v0.3.3.6-issue36'
STAGE = OUT / 'D-canonical-roles'
v = json.loads((STAGE / 'validation.json').read_text())
p = json.loads((STAGE / 'preservation.json').read_text())
sha = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
assert v['source_sha256'] == p['source_sha256'] == sha
assert v['behavior_total'] == 854 and v['static']['passed'] == 21
assert v['errors'] == 0 and v['gta_testing'] == 'NOT RUN'
code_tip = 'e4860da7accc386ab4b0d36c9c85b0cc4c55f7b7'
assert SOURCE.read_bytes() == subprocess.check_output(['git', 'show', code_tip+':'+str(SOURCE.relative_to(ROOT))], cwd=ROOT)
commits = subprocess.check_output(['git', 'log', '--reverse', '--format=%H %s', BASE+'..'+code_tip], cwd=ROOT, text=True).strip()
src = str(SOURCE.relative_to(ROOT))
diff = subprocess.check_output(['git', 'diff', BASE, code_tip, '--', src], cwd=ROOT)
(OUT / 'source.diff').write_bytes(diff)
with tempfile.TemporaryDirectory(prefix='gh36-diff-review-') as temp:
    target = Path(temp) / src
    target.parent.mkdir(parents=True)
    target.write_bytes(subprocess.check_output(['git', 'show', BASE+':'+src], cwd=ROOT))
    subprocess.run(['git', 'apply', str(OUT / 'source.diff')], cwd=temp, check=True)
    assert target.read_bytes() == SOURCE.read_bytes()

old, new = members(baseline()), members(SOURCE.read_text())
excerpts = '# Issue #36 before/after source excerpts\n\nFull exact diff: source.diff. Base: '+BASE+'.\n'
for title, method, needle in [
    ('A footer layout', 'DrawMenu', 'string[] wrapped'),
    ('B work-stop labels', 'GetMenuItems', 'case MenuPage.MechanicWalkingPaths:'),
    ('C explicit warehouse GPS', 'SetSupplierPickupWaypoint', None),
    ('D canonical staff route', 'ExecuteReceptionistManagement', None),
]:
    excerpts += '\n## '+title+' — '+method+'\n'
    for label, content in [('Before', old[method]), ('After', new[method])]:
        lines = content.splitlines()
        if needle:
            at = next(i for i, line in enumerate(lines) if needle in line)
            lines = lines[max(0, at-1):at+10]
        excerpts += '\n'+label+':\n\n```csharp\n'+'\n'.join(lines)+'\n```\n'
(OUT / 'SOURCE_EXCERPTS.md').write_text(excerpts)

guide = f'''# Glass Heroes Business v0.3.3.6 — Issue #36 maintenance review

EXPERIMENTAL — PENDING OWNER ACCEPTANCE. Report date: October 24, 2026. GTA testing: NOT RUN. Windows execution: NOT RUN. This is a review candidate, not an accepted release or a crash fix. The stack-overflow cause remains UNATTRIBUTED.

## Identity and approvals

Branch: maintenance/v0.3.3.6-issue36

Pinned RC1 base: {BASE}

Final source/four-fix checkpoint: {code_tip}

Source SHA256: {sha}

The package's FINAL_HANDOFF.md and FINAL_HANDOFF.json record the exact final branch tip including the subsequent documentation commit. This tracked guide names the code checkpoint to avoid a self-referential documentation commit hash.

Gate 1: owner YES in the current conversation, before implementation. APPROVALS.md supersedes the historical awaiting-approval wording in APPROVED_PLAN.md. Gate 2: PENDING owner YES/NO after reviewing this candidate. No acceptance, merge, deployment, tag or v0.3.4 integration has been performed.

Individual fix commits, in order:

```
{commits}
```

## A — footer descriptions

The previous fixed-character split could overflow proportional-font text, leave a single trailing word and crowd the control hints. WrapMenuDescription now uses measured font widths and a bounded cached LayoutMenuDescription/FittingMenuWords helper. It balances two fitting word lines, normalizes whitespace/newline tokens and uses a word-boundary ellipsis if two lines are insufficient. Oversized input and tokens are bounded safely. DrawMenu preserves the 420-pixel panel, 119-pixel footer, font, colours, rows and hint/counter positions. Description baselines are spaced 23 design pixels; single-line help is centered vertically. The generic Back hint is shorter.

Placement uses the existing SHVDN 1280x720 reference space at MenuScale; no resolution query or wrapping calculations occur for an unchanged cached description. Cache capacity is 128; input is bounded to 1024 characters and 96 words; fitting uses binary search. GTA font/native rendering, safe-zone, 720p/1080p and overlay readability still need owner observation. No screenshots were inspected by Codex.

## B — mechanic terminology

Shared Entry Corridor means the existing physical entry path, also used by the existing optional car pickup/return walking toggle. Shared Work Stops means the five shared workstation-access nodes. Workstation Approaches identifies each separate mechanic final approach. Selections by Mechanic / To Workstation / From Workstation identify per-worker masks selecting those shared work stops. These do not rename or move the underlying saved data.

MechanicRouteCounts reports selected, physically saved and unset counts. MechanicRoutePreview displays the existing ascending outbound and descending return order, with the approach/target; DIRECT mode marks custom selections inactive. MechanicRouteHeading fits direction-first text without hiding direction. Preview is read-only and never calls BuildRoute, Save or pathfinding. The 16 route cases include 6144 mask/saved-mask/direction/worker combinations; these combinations are inside those cases, not 6144 additional tests. AI, traversal order, geometry, bays and waypoint masks are unchanged.

## C — READY physical wholesale locator

After startup, SyncWholesalePickupBlip runs in the existing active, unpaused one-second supplier refresh. An actual wholesale order with stock, finished ETA, stage 7, a valid existing warehouse point and no carried/loaded/unloaded cargo owns one long-range yellow Package blip named Glass Heroes — Wholesale Parts Pickup. It remains useful outside the prop streaming radius; no actor, model or distance scan is required. Transit and manager deliveries do not receive this locator.

Active Deliveries adds Set Warehouse GPS at row 4, before Back. Existing action rows 0–3 keep their behavior. GPS is explicit and READY-only: it enables the owned blip's ShowRoute without setting, deleting or restoring the player's waypoint. A later changed player waypoint yields the owned route at the next active check. Ordering/READY do not automatically seize GPS. Existing post-collection return-to-shop navigation is unchanged.

Three creation attempts are spaced five active seconds apart, followed by a 30-second cooldown and recovery window. Logging/setup messages are rate limited to 30 active seconds. Valid handles are reused; partial configuration is cleaned. Failed deletion retains ownership and blocks duplicates, with at most two delete attempts per synchronization. Successful collection/loading, cancellation, recovery, abort and invalid/completed order state clean only the owned locator. The READY saved order reconstructs the locator after the existing startup gate; GPS choice is session-only and starts off after reload. The native engine can still fail to create/delete a blip; adapters test these failures but do not prove GTA cleanup behavior.

Collection remains entirely player-operated. No new function bills, stocks inventory, collects cargo, moves staff or saves business data. Existing depot pickup point, world marker, physical props and accounting remain unchanged. Missing/nonfinite locations get setup help rather than invented coordinates.

## D — one staff-route editor per role

Placement Settings > Walking Routes, role-management and Employee Walking Paths shortcuts now open the existing ReceptionWalkingPaths or ManagerWalkingPaths canonical page. Each retains walking mode, five saved stops, final approach and Back. Redundant visible wrapper steps are bypassed; internal enum aliases are retained for compatibility. Existing OpenPlacementMenu context records the real caller/row and the child-list return, replacing stale menu context. NavigateBack and PointEditor themselves remain unchanged.

Receptionist Walk Route and Manager Walk Route titles/help identify the same settings from every shortcut. Mode guards, Save/Clear actions and coordinates/headings are unchanged. The navigation suite retains its original 180 checks with only approved destinations/row expectations updated and adds 28 cases for role origins, all stops/approaches, explicit save, modes, stale reopen context, keyboard/controller/mixed input and pause.

## Exact source-level scope and data compatibility

The only gameplay source changed is {src}. The Visual Studio project still links this exact install source; project/solution and original v2 sources are byte-identical to RC1. SOURCE_EXCERPTS.md and source.diff contain before/after evidence and the complete unified source diff; applying it to the pinned base reproduced the source byte-for-byte.

Changed existing members:

{', '.join(p['changed_methods'])}.

New members:

{', '.join(p['new_members'])}.

The guard confirms {p['unchanged_members']} existing members and all other existing root declarations/defaults are unchanged. Only approved presentation caches and owned-locator session fields are added. All GetValue/SetValue expressions are identical. Save.ini and Settings.ini have zero new keys, defaults, schema changes, migrations, coordinate rewrites or mask resets. No existing user save/settings file was read or modified. Synthetic data is used only by tests. Financial transactions, inventory transfers, pending-order identity, invoicing, employee AI, route geometry, active-unpaused clocks, constructor/startup defaults and controller arbitration remain unchanged. The guard also checks every allowed locator hook is outside the protected transaction logic.

## Recorded compilation and regression results

Actual clean Release compilation: PASS, 0 errors, two inherited warnings CS0169 and CS0649. SDK 8.0.426; .NET Framework 4.8; C# 7.3; official ScriptHookVDotNet3 package pinned to 3.6.0. This is real product compilation, not a stub-only build. Warning lines are in D-canonical-roles/build.txt.

Behavior/navigation: 854/854 PASS, compared with independently reproduced RC1 baseline 762/762. Static: 21/21 PASS. Approved-scope/preservation guard: PASS. New cases: footer36 15, routes36 16, locator36 33, plus 28 navigation cases. Final navigation: 208/208. Inherited issue/logger/migration checks: 85. Other feature suites: {', '.join(k+' '+str(x['total']) for k,x in v['feature_suites'].items() if not k.endswith('36'))}.

Behavior tests extract actual production members/case arms and compile/run them against deterministic adapters on .NET 8. They cover accounting/clock/input/native-failure behavior but do not execute GTA natives or prove real GPS rendering, controller hardware, collision, MLO navigation or crash stability. Static checks and preservation are separate from the 854 behavior count. Every stage's actual logs are retained; D-canonical-roles/validation.json is the final source's result. No identical gameplay suite rerun was needed for documentation-only changes.

Reproduce from the Git branch with Python 3 and .NET SDK 8 available (configured cloud shell: source /workspace/.cloud-setup/env.sh):

```
python verification/maintenance36/verify.py --stage reviewer-rerun --extra footer36 routes36 locator36
```

The complete verifier requires the pinned Git base and historical validation/protected files. The review ZIP includes source/project, tests, historical baseline text, evidence and diffs; clean product rebuild does not require Git history. Product build:

```
dotnet build experimental/shvdn3/VisualStudio/GlassHeroesBusiness.V3.Experimental.csproj -c Release
```

A .NET 8 SDK supplies build tooling and test execution; the product remains net48. The package extraction receipt records a separate clean product build from the archive. No compiled DLL is supplied for simultaneous installation with the source script.

## Review package, installation and rollback

This package is prepared for code review before acceptance. It has one active GlassHeroesBusiness.3.cs under experimental/shvdn3/Install/scripts and the matching linked Visual Studio project, documentation, tests and actual logs. No Save.ini, Settings.ini or alternate runtime script/DLL is included. Historical source fixtures are text, not installable scripts. Do not copy the entire archive into GTA.

No installation is performed or authorized by this report. If the owner later accepts and authorizes testing: close GTA; back up the exact currently tested source and complete matching Glass Heroes data/settings directory together; use a disposable copy of the data for first tests; install only the reviewed .3.cs using the owner's existing SHVDN3 setup. Ensure only one Business Manager generation runs; keep older .2.cs scripts or compiled Business Manager DLLs outside the active scripts folder. Do not overwrite saved data/settings with examples. Record archive/source hashes and runtime/dependency/mod-stack versions.

Rollback with GTA closed: restore the previously tested source and its complete matching data/settings backup together. Keep copies of both pre-test and post-test data for investigation; never reset or merge sole live saves to roll back code. Source does not add a migration, but gameplay transactions during testing can change balances/inventory normally. Report any missing shipment, double invoice, changed coordinate or pause/offline advancement immediately with copies and logs.

## Integration ownership

The separate v0.3.4 instance owns integration only after explicit owner gate 2 YES. This branch contains no v0.3.4 features. Do not overwrite its newer monolithic source with this RC1-based file. Review the four deltas against its current source, resolve shared menu/wholesale overlaps deliberately, rerun its complete build/regressions and obtain its review. Test fixture adapter changes preserve prior assertions and must accompany the corresponding hooks if needed. Do not merge overlapping historical QA/issue PR sets independently. This instance has not changed or merged main, RC1, QA R2, v0.3.4 or live GTA.
'''
(OUT / 'DEVELOPER_GUIDE.md').write_text(guide)

qa = '''# Issue #36 GTA acceptance checklist — ALL NOT RUN

Gate 2 remains PENDING YES/NO. This checklist describes later owner testing, not work Codex performed. Preserve exact matching source plus full data/settings backups before any authorized test. Record result, game/runtime/mod versions, source SHA, screenshots and logs per case. Crash cause remains UNATTRIBUTED.

1. Cold launch with the heavy mod stack; observe existing startup delay, closed/open shop and F2 behavior. Reload the script and restart GTA with copies of an existing saved business. Check money, stock, invoices, active jobs, employee modes, waypoint masks/coordinates and settings before/after.
2. At 720p and 1080p (and the owner's resolution/safe-zone), visit long/short help, Back, route previews and Active Deliveries. Confirm two readable word lines, honest ellipsis, no footer/hint/counter overlap, preserved panel appearance and HUD/other-mod overlay readability. Inspect actual screenshots; adapter geometry is not visual qualification.
3. Mechanic labels: distinguish Shared Entry Corridor, car pickup/return toggle, Shared Work Stops, workstation approaches and per-worker/direction selections. Use sparse/unsaved masks and long worker names. Confirm selected/saved/unset counts and DIRECT inactive status. Opening previews must not save or change anything.
4. With copied custom layouts, run existing pickup, workstation and return paths for all mechanics/directions; compare physical traversal and bays with RC1. No geometry change should be observed. Change only explicit settings on disposable data. Verify pause/menu/closed-game time does not advance business progress, wages, supplier ETA or remote income.
5. Enter both staff-route editors from Placement Settings, role management and Employee Walking Paths. Visit each of five stops and final approach; use cancel/Back, explicit Save/Clear and mode guard with shop open/closed. Confirm Back restores caller and row. Reopen F2 after a different caller; verify stale context does not hijack Back.
6. Repeat navigation with keyboard-only, controller-only and mixed input, including held/repeated inputs, menu paging, pause/resume and controller Back. Confirm one action per press and no saves from passive navigation.
7. Order physical wholesale stock and inspect transit: no READY pickup blip or automatic GPS. On READY, confirm one named long-range package blip when far away, near the warehouse and after prop streaming transitions. Manager supplier deliveries must not use this blip. Collection must still require the player.
8. Select Set Warehouse GPS explicitly for a READY order, both with and without an existing player waypoint. Verify owned route appears without changing that waypoint. Choose a different waypoint and confirm the owned route yields within the next active second. Verify GPS cannot activate for transit, empty, carried, loaded or completed orders.
9. Collect/load/unload/deposit wholesale stock with the existing workflow. Confirm owned locator cleanup, normal existing home navigation, exactly one inventory transfer/invoice and unchanged finances. Cancel successfully, recover, reload READY/transit/loaded orders and abort/reload scripts: no stale/duplicate blips, no automatic GPS reconstruction, no lost cargo or duplicate payment. Failed cancellation must preserve the order/locator. Native-failure cases may require controlled diagnostic setup; do not corrupt live data to force them.
10. Check missing point/error help where safely reproducible on disposable data. Verify no invented pickup destination, repeated message spam or per-frame performance regression. Idle/away/streaming stress-test and collect crash logs if any. A successful session does not establish the historical stack-overflow cause.

For any failure, retain the exact source/package hash, pre/post disposable saves/settings, reproduction sequence, time to failure, screenshots, SHVDN/business logs and crash/minidump evidence when available. Do not attribute a crash without evidence.
'''
(OUT / 'GTA_TEST_CHECKLIST.md').write_text(qa)

lines = (STAGE / 'navigation-tests.txt').read_text().splitlines()
matrix = '# Issue #36 production navigation inventory\n\nGenerated from the actual final navigation run. Destinations/callers and row labels are source-derived; in-game navigation remains NOT RUN.\n\n'
footers = '# Issue #36 production row descriptions\n\nSource-derived descriptions emitted by the final navigation harness. Font rendering and visual readability in GTA: NOT RUN.\n\n'
for line in lines:
    if line.startswith('MATRIX|'):
        _, page, title, parent, selection, rows = line.split('|', 5)
        matrix += f'## {page} — {title}\n\nHarness return: {parent}, row {selection}.\n\n'+ '\n'.join(f'- {i}: {row}' for i, row in enumerate(rows.split('; ')))+'\n\n'
    elif line.startswith('FOOTER|'):
        _, page, row, label, description = line.split('|', 4)
        footers += f'- {page} [{row}] {label}: {description}\n'
(OUT / 'NAVIGATION_MATRIX.md').write_text(matrix)
(OUT / 'MENU_FOOTERS.md').write_text(footers)
(OUT / 'CHANGELOG.md').write_text('''# v0.3.3.6 Issue #36 maintenance changelog — EXPERIMENTAL

October 24, 2026. Gate 1 YES; gate 2 PENDING. GTA testing NOT RUN.

- A: measured/cached two-line footer wrapping, balanced spacing and safe ellipsis within existing menu appearance.
- B: distinguish shared entry/work stops, approaches and mechanic/direction selections; read-only counts/preview, unchanged AI/geometry/masks.
- C: owned long-range READY physical-wholesale pickup blip, explicit GPS that yields to changed player waypoint, bounded retries and lifecycle cleanup. Player still collects; original accounting and active clocks preserved.
- D: common receptionist/manager route pages from every shortcut, retaining saved settings, explicit edits and contextual Back.
- Actual clean Release: 0 errors / two inherited warnings; 854/854 behavior/navigation and 21/21 static checks PASS. No schema/default changes or data rewrite. No claim of GTA qualification or crash resolution.
''')
(OUT / 'INTEGRATION_HANDOFF.md').write_text(f'''# Issue #36 integration handoff — PENDING OWNER ACCEPTANCE

Exact branch: maintenance/v0.3.3.6-issue36

Base: {BASE}

Source/four-fix tip: {code_tip}

Final branch tip including documents: see FINAL_HANDOFF.md/.json in the review package. The external receipt is written after the final documentation commit so its SHA is exact.

{commits}

Product source: {src}

SHA256: {sha}

Complete changed-file list, per-commit changed files, final documentation SHA and build results accompany the exact-tip final handoff. source.diff is the complete product delta; SOURCE_EXCERPTS.md and COMMIT_A/B/C/D.md explain each fix. Final evidence: D-canonical-roles/. 854/854 behavior/navigation, 21/21 static, preservation PASS, actual net48 Release 0 errors/two old warnings. GTA and Windows execution NOT RUN. No crash attribution.

Save compatibility: zero GetValue/SetValue expression or schema/key/default changes; no Save.ini/Settings.ini read or modified; no finance/clock/AI/coordinate/mask change. Session-only footer cache and owned map locator. Preserve exact source plus full matching data backups for any future test/rollback.

Gate 1 YES. Gate 2 PENDING owner YES/NO. Do not integrate or accept until gate 2 YES. The separate v0.3.4 instance owns the later port/review and must handle shared source/menu overlaps deliberately; do not replace its source wholesale. This instance performs no v0.3.4 merge or modification. Original RC1, main, QA R2 and live game untouched. Review-only ZIP is not a published release.
''')
(OUT / 'PR_BODY.md').write_text('''Fixes #36

Draft maintenance candidate based on pinned v0.3.3.6 RC1, isolated on maintenance/v0.3.3.6-issue36. Four separate fix commits improve footer layout, distinguish mechanic route terms with read-only summaries, add an owned READY wholesale pickup blip/explicit GPS, and route staff shortcuts into their existing canonical editors.

Actual net48/C#7.3 Release against SHVDN3 3.6.0: 0 errors, two inherited warnings. 854/854 behavior/navigation, 21/21 static and preservation guard PASS. Save/settings expressions, mechanic AI/geometry, finances and active-unpaused clocks are unchanged. GTA and Windows execution NOT RUN; crash cause UNATTRIBUTED.

See docs/features/v0.3.3.6-issue36/DEVELOPER_GUIDE.md, source.diff, actual logs, Word/PDF, acceptance checklist and exact-tip package handoff. Owner implementation gate YES; final acceptance gate PENDING YES/NO. Keep draft. Do not merge into RC1, main or v0.3.4. The separate v0.3.4 instance owns integration after owner acceptance. No live installation or business data touched.
''')

pdfmetrics.registerFont(TTFont('ReportText', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
pdfmetrics.registerFont(TTFont('ReportBold', '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'))
pdfmetrics.registerFont(TTFont('ReportMono', '/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf'))
styles = getSampleStyleSheet()
for name in ['BodyText', 'Normal', 'Title', 'Heading2']:
    styles[name].fontName = 'ReportBold' if name in ['Title', 'Heading2'] else 'ReportText'
styles.add(ParagraphStyle(name='CodeSmall', fontName='ReportMono', fontSize=7, leading=9))
doc = Document()
doc.styles['Normal'].font.size = Pt(10)
flow, code = [], False
for line in (guide+'\n'+qa+'\n'+excerpts).splitlines():
    if line.startswith('```'):
        code = not code
        continue
    if code:
        run = doc.add_paragraph().add_run(line)
        run.font.name, run.font.size = 'Courier New', Pt(8)
        for part in textwrap.wrap(line, 94, replace_whitespace=False, drop_whitespace=False) or ['']:
            flow.append(Preformatted(part, styles['CodeSmall']))
    elif not line:
        flow.append(Spacer(1, 5))
    elif line.startswith('#'):
        n = len(line)-len(line.lstrip('#'))
        title = line[n:].strip()
        doc.add_heading(title, level=min(n-1, 3))
        flow.append(Paragraph(escape(title), styles['Title' if n == 1 else 'Heading2']))
    else:
        doc.add_paragraph(line)
        flow.append(Paragraph(escape(line), styles['BodyText']))
stem = 'GlassHeroesBusiness_v0.3.3.6_Issue36_Maintenance_Developer_Guide'
doc.save(OUT / (stem+'.docx'))
def footer(canvas, document):
    canvas.setFont('ReportText', 8)
    canvas.drawString(42, 22, 'Issue #36 | EXPERIMENTAL | acceptance PENDING | GTA NOT RUN')
    canvas.drawRightString(570, 22, str(document.page))
pdf = OUT / (stem+'.pdf')
SimpleDocTemplate(str(pdf), pagesize=(612, 792), leftMargin=42, rightMargin=42, topMargin=40, bottomMargin=38).build(flow, onFirstPage=footer, onLaterPages=footer)
reader = PdfReader(pdf)
for content in ['\n'.join(x.text for x in Document(OUT / (stem+'.docx')).paragraphs), '\n'.join(x.extract_text() for x in reader.pages)]:
    for term in [sha, code_tip, '854/854', '21/21', 'Save.ini', 'Settings.ini', 'PENDING', 'NOT RUN']:
        assert term in content, term
(OUT / 'SOURCE_SHA256.txt').write_text(sha+'  '+src+'\n')
(OUT / 'report-receipt.json').write_text(json.dumps({'code_commit':code_tip, 'source_sha256':sha, 'diff_reproduced':True, 'word_pdf_text_verified':True, 'pdf_pages':len(reader.pages), 'visual_qa':'NOT RUN', 'gta_testing':'NOT RUN', 'owner_acceptance':'PENDING YES/NO'}, indent=2)+'\n')
print(json.dumps({'source_sha256':sha, 'diff_reproduced':True, 'pdf_pages':len(reader.pages), 'word_pdf_text_verified':True}))
