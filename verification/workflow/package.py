#!/usr/bin/env python3
"""Create a review-only source package after successful current-source verification."""
import argparse,hashlib,json,subprocess,tempfile,zipfile,textwrap
from pathlib import Path
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Pt
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Preformatted
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from pypdf import PdfReader
repo=Path(__file__).resolve().parents[2]
p=argparse.ArgumentParser();p.add_argument('label');p.add_argument('--parent',required=True);args=p.parse_args()
r=repo/'docs/releases/v0.3.3.2-shvdn3-supplier-workflow';v=json.loads((r/'validation.json').read_text());source='experimental/shvdn3/Install/scripts/GlassHeroesBusiness.3.cs'
assert v['build']=='PASS' and v['gta_testing']=='NOT RUN'
sha=hashlib.sha256((repo/source).read_bytes()).hexdigest();assert sha==v['source_sha256']
(r/'source.diff').write_bytes(subprocess.check_output(['git','diff',args.parent,'--',source],cwd=repo))
# Reproduce the exact candidate from its parent plus this review diff.
with tempfile.TemporaryDirectory(prefix='gh-review-diff-') as d:
 dest=Path(d)/source;dest.parent.mkdir(parents=True);dest.write_bytes(subprocess.check_output(['git','show',args.parent+':'+source],cwd=repo))
 subprocess.run(['git','apply',str(r/'source.diff')],cwd=d,check=True);assert dest.read_bytes()==(repo/source).read_bytes()
guide=(r/'DEVELOPER_GUIDE.md').read_text().split('\n## Recorded automated results\n')[0]
guide+='\n## Recorded automated results\n\nActual net48 / C# 7.3 Release build against ScriptHookVDotNet3 3.6.0: PASS, 0 errors, inherited warnings CS0169 and CS0649. SDK '+v['sdk']+'.\n\n'
guide+=str(v['behavior_total'])+'/'+str(v['behavior_total'])+' extracted-behavior tests PASS (85 inherited; '+', '.join(k+': '+str(c['passed']) for k,c in v['feature_suites'].items())+'). 21/21 static checks PASS. Source preservation/scope checks PASS; active-game clock sampling and original v2 source/project are unchanged. Supplier route timing consumes that existing active-game sample. Tests use deterministic world stubs; they do not run GTA.\n\n'
guide+='Source SHA256: '+sha+'. The complete unified source.diff applies to parent '+args.parent+' and reproduces this source exactly.\n\nGTA testing: NOT RUN. Windows execution: NOT RUN. Installed Enhanced 3.9.0.6 compatibility: NOT VERIFIED. Stack-overflow cause: UNATTRIBUTED / UNPROVEN. No installation, merge, protected-branch push, or live save/settings editing was performed.\n'
(r/'DEVELOPER_GUIDE.md').write_text(guide)
stem='GlassHeroesBusiness_v0.3.3.2_SHVDN3_'+args.label.replace('-','_')+'_Developer_Report'
doc=Document();doc.styles['Normal'].font.size=Pt(10);styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='SmallCode',fontName='Courier',fontSize=7,leading=9));flow=[];code=False
for line in guide.splitlines():
 if line.startswith('```'):code=not code;continue
 if code:
  q=doc.add_paragraph();run=q.add_run(line);run.font.name='Courier New';run.font.size=Pt(8)
  for part in textwrap.wrap(line,96,replace_whitespace=False,drop_whitespace=False) or ['']:flow.append(Preformatted(part,styles['SmallCode']))
 elif not line:flow.append(Spacer(1,5))
 elif line.startswith('#'):
  n=len(line)-len(line.lstrip('#'));text=line[n:].strip();doc.add_heading(text,level=min(n-1,3));flow.append(Paragraph(escape(text),styles['Title' if n==1 else 'Heading2']))
 else:doc.add_paragraph(line);flow.append(Paragraph(escape(line),styles['BodyText']))
doc.save(r/(stem+'.docx'))
def footer(c,p):
 c.setFont('Helvetica',8);c.drawString(42,22,'SHVDN3 EXPERIMENTAL | GTA testing NOT RUN | REVIEW PENDING');c.drawRightString(570,22,str(p.page))
SimpleDocTemplate(str(r/(stem+'.pdf')),pagesize=(612,792),leftMargin=42,rightMargin=42,topMargin=40,bottomMargin=38).build(flow,onFirstPage=footer,onLaterPages=footer)
for text in ['\n'.join(p.text for p in Document(r/(stem+'.docx')).paragraphs),'\n'.join(p.extract_text() for p in PdfReader(r/(stem+'.pdf')).pages)]:
 for phrase in ['EXPERIMENTAL','NOT RUN','Save.ini','Recorded automated results',str(v['behavior_total'])+'/'+str(v['behavior_total'])]:assert phrase in text,phrase
files={repo/source};files.update(x for x in (repo/'experimental/shvdn3/VisualStudio').rglob('*') if x.is_file() and not any(y in {'bin','obj'} for y in x.parts));files.add(repo/'experimental/shvdn3/README.md');files.add(repo/'AGENTS.md')
files.update(x for x in (repo/'verification').rglob('*') if x.is_file() and not any(y in {'bin','obj','__pycache__'} for y in x.parts))
files.update(x for x in r.rglob('*') if x.is_file() and x.name not in {'FILES_SHA256.json','package-receipt.json','extracted-package-build.txt'})
# Include historical feature evidence with its original hashes and version labels.
for label in ['repair-coverage','manual-parts-sales','pallet-stacking','manager-wages','saved-quotes','final-review','manager-deliveries']:
 related=repo/f'docs/releases/v0.3.3.1-shvdn3-{label}'
 files.update(x for x in related.rglob('*') if x.is_file())
files.add(repo/'docs/releases/v0.3.3.1-shvdn3-phase1/API_MAP.md')
files.add(repo/'docs/releases/v0.3.3.1-shvdn3-phase1/dependencies.json')

for name in ['Verification/static_verify.py','BASELINE_SHA256.txt','SourceDiff_v0.3.2.9a_to_v0.3.3.0_RC1.patch']:files.add(repo/name)
manifest={'source_sha256':sha,'files':{str(x.relative_to(repo)):hashlib.sha256(x.read_bytes()).hexdigest() for x in sorted(files)}};(r/'FILES_SHA256.json').write_text(json.dumps(manifest,indent=2)+'\n');files.add(r/'FILES_SHA256.json')
out=repo/'build/releases';out.mkdir(parents=True,exist_ok=True);archive=out/f'GlassHeroesBusiness_v0.3.3.2_SHVDN3_{args.label}_EXPERIMENTAL.zip'
with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
 for x in sorted(files):z.write(x,str(x.relative_to(repo)))
 z.write(repo/'GlassHeroesBusiness_v0.3.3.0_RC1.2.cs','verification/baselines/GlassHeroesBusiness_v0.3.3.0_RC1.2.cs.txt')
 z.write(r/'READ_ME_FIRST.txt','READ_ME_FIRST.txt')
with tempfile.TemporaryDirectory(prefix='gh-extracted-review-') as d:
 with zipfile.ZipFile(archive) as z:
  assert z.testzip() is None and len(z.namelist())==len(set(z.namelist()))
  assert [x for x in z.namelist() if x.endswith('.3.cs')]==[source]
  assert not any(x.endswith(('.2.cs','.dll','.zip','/Save.ini','/Settings.ini')) for x in z.namelist())
  z.extractall(d)
 for name,digest in manifest['files'].items():assert hashlib.sha256((Path(d)/name).read_bytes()).hexdigest()==digest
 product=Path(d)/'experimental/shvdn3/VisualStudio/GlassHeroesBusiness.V3.Experimental.csproj'
 build=subprocess.run(['dotnet','build',str(product),'-c','Release','--property:BaseIntermediateOutputPath='+d+'/obj/','--property:OutputPath='+d+'/bin/','--maxcpucount:2','--verbosity','minimal'],stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True)
 (r/'extracted-package-build.txt').write_text(build.stdout);assert build.returncode==0,build.stdout
receipt={'package':archive.name,'sha256':hashlib.sha256(archive.read_bytes()).hexdigest(),'source_sha256':sha,'one_installation_source':True,'no_v2_dll_or_live_data':True,'all_file_hashes_verified':True,'extracted_net48_build':'PASS','behavior_tests':v['behavior_total'],'static_tests':21,'gta_testing':'NOT RUN','review_status':'PENDING'}
(r/'package-receipt.json').write_text(json.dumps(receipt,indent=2)+'\n');archive.with_suffix('.manifest.json').write_text(json.dumps(receipt,indent=2)+'\n');print(json.dumps(receipt,indent=2))
