#!/usr/bin/env python3
"""Generate review documents and hashes only. Does not publish/package/install a release."""
import hashlib,json,subprocess,tempfile,textwrap
from pathlib import Path
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Pt
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Preformatted
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from pypdf import PdfReader
repo=Path(__file__).resolve().parents[2];r=repo/'docs/features/v0.3.3.6';v=json.loads((r/'validation.json').read_text());src='experimental/shvdn3/Install/scripts/GlassHeroesBusiness.3.cs';sha=hashlib.sha256((repo/src).read_bytes()).hexdigest();assert v['source_sha256']==sha and v['build']=='PASS' and v['gta_testing']=='NOT RUN'
(r/'source.diff').write_bytes(subprocess.check_output(['git','diff','22bff09','--',src],cwd=repo))
with tempfile.TemporaryDirectory(prefix='gh336-review-diff-') as d:
 p=Path(d)/src;p.parent.mkdir(parents=True);p.write_bytes(subprocess.check_output(['git','show','22bff09:'+src],cwd=repo));subprocess.run(['git','apply',str(r/'source.diff')],cwd=d,check=True);assert p.read_bytes()==(repo/src).read_bytes()
guide=(r/'DEVELOPER_GUIDE.md').read_text().split('\n## Recorded automated results\n')[0]
guide+='\n## Recorded automated results\n\nGenuine SHVDN3 3.6.0 net48 / C#7.3 Release: PASS, 0 errors; inherited warnings CS0169 and CS0649. SDK '+v['sdk']+'.\n\n'
guide+=str(v['behavior_total'])+'/'+str(v['behavior_total'])+' behavior tests PASS: 85 inherited issue/logger/migration checks plus '+', '.join(k+' '+str(x['total']) for k,x in v['feature_suites'].items())+'; navigation '+str(v['navigation_tests']['total'])+'. Static 21/21 PASS; focused scope/save/v2/clock/settlement/transfer-helper preservation PASS. The new navigation tests use production members/case arms with deterministic adapters, not GTA.\n\n'
guide+='Candidate source SHA256: '+sha+'. Full diff verified to reproduce this source from 22bff09.\n\nGTA testing NOT RUN. Windows execution NOT RUN. Enhanced runtime compatibility NOT VERIFIED. Crash cause UNATTRIBUTED. No release archive, runtime deployment, save modification or merge was produced by this task.\n'
(r/'DEVELOPER_GUIDE.md').write_text(guide)
styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='CodeSmall',fontName='Courier',fontSize=7,leading=9));doc=Document();doc.styles['Normal'].font.size=Pt(10);flow=[];code=False
for line in (guide+'\n\n'+(r/'GTA_TEST_CHECKLIST.md').read_text()+'\n\n'+(r/'NAVIGATION_MATRIX.md').read_text()).splitlines():
 if line.startswith('```'):code=not code;continue
 if code:
  q=doc.add_paragraph();run=q.add_run(line);run.font.name='Courier New';run.font.size=Pt(8)
  for part in textwrap.wrap(line,96,replace_whitespace=False,drop_whitespace=False) or ['']:flow.append(Preformatted(part,styles['CodeSmall']))
 elif not line:flow.append(Spacer(1,5))
 elif line.startswith('#'):
  n=len(line)-len(line.lstrip('#'));t=line[n:].strip();doc.add_heading(t,level=min(n-1,3));flow.append(Paragraph(escape(t),styles['Title' if n==1 else 'Heading2']))
 else:doc.add_paragraph(line);flow.append(Paragraph(escape(line),styles['BodyText']))
stem='GlassHeroesBusiness_v0.3.3.6_Development_Report';doc.save(r/(stem+'.docx'))
def footer(c,p):
 c.setFont('Helvetica',8);c.drawString(42,22,'Issues #33/#34 DEVELOPMENT REVIEW | EXPERIMENTAL | GTA validation PENDING');c.drawRightString(570,22,str(p.page))
SimpleDocTemplate(str(r/(stem+'.pdf')),pagesize=(612,792),leftMargin=42,rightMargin=42,topMargin=40,bottomMargin=38).build(flow,onFirstPage=footer,onLaterPages=footer)
for text in ['\n'.join(p.text for p in Document(r/(stem+'.docx')).paragraphs),'\n'.join(p.extract_text() for p in PdfReader(r/(stem+'.pdf')).pages)]:
 for term in ['EXPERIMENTAL','NOT RUN','Save.ini',str(v['behavior_total'])+'/'+str(v['behavior_total']),sha]:assert term in text,term
(r/'SOURCE_SHA256.txt').write_text(sha+'  '+src+'\n')
files=[repo/src,repo/'experimental/shvdn3/VisualStudio/GlassHeroesBusiness.V3.Experimental.csproj']+sorted(p for p in r.iterdir() if p.is_file() and p.name not in ['FILES_SHA256.json','report-receipt.json'])
(r/'FILES_SHA256.json').write_text(json.dumps({'source_sha256':sha,'files':{str(p.relative_to(repo)):hashlib.sha256(p.read_bytes()).hexdigest() for p in files}},indent=2)+'\n')
(r/'report-receipt.json').write_text(json.dumps({'source_sha256':sha,'source_diff_reproduced':True,'word_pdf_verified':True,'behavior_tests':v['behavior_total'],'static_checks':21,'gta_testing':'NOT RUN','release_published':False,'new_release_zip':False,'live_installation_modified':False},indent=2)+'\n')
print('Review documents verified; no release archive created. Source '+sha)
