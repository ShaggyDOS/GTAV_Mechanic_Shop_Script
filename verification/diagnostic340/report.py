#!/usr/bin/env python3
"""Build the RC2 diagnostic review documents without deployment or acceptance."""
import hashlib,json,subprocess,sys,tempfile,textwrap
from pathlib import Path
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Pt
from pypdf import PdfReader
from reportlab.lib.styles import ParagraphStyle,getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph,Preformatted,SimpleDocTemplate,Spacer
ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'verification/maintenance36'))
from source_tools import SOURCE,members
BASE='e312d8c9500f4a990eb993b82e8e76b55711d3de'
CODE='3b2ef09a96dda9f6d0325f08d29bc986167abe68'
OUT=ROOT/'docs/features/v0.3.4-rc2-diagnostic'
v=json.loads((OUT/'final-validation/validation.json').read_text())
sha=hashlib.sha256(SOURCE.read_bytes()).hexdigest()
assert sha==v['source_sha256'] and v['behavior']['total']==985 and v['build']['errors']==0
assert SOURCE.read_bytes()==subprocess.check_output(['git','show',CODE+':'+str(SOURCE.relative_to(ROOT))],cwd=ROOT)
diff=subprocess.check_output(['git','diff',BASE,CODE,'--',str(SOURCE.relative_to(ROOT))],cwd=ROOT)
(OUT/'source.diff').write_bytes(diff)
with tempfile.TemporaryDirectory(prefix='gh-rc2-diff-') as directory:
    target=Path(directory)/SOURCE.relative_to(ROOT);target.parent.mkdir(parents=True)
    old=subprocess.check_output(['git','show',BASE+':'+str(SOURCE.relative_to(ROOT))],cwd=ROOT)
    target.write_bytes(old)
    subprocess.run(['git','apply',str(OUT/'source.diff')],cwd=directory,check=True)
    assert target.read_bytes()==SOURCE.read_bytes()
before,after=members(old.decode()),members(SOURCE.read_text())
excerpts='# RC2 before/after source excerpts\n\nFull exact unified diff: source.diff. Gameplay bodies are preserved inside new diagnostic scopes.\n'
for name in ['BeginEnterVehicle','BeginVehicleDrive']:
    excerpts+='\n## '+name+'\n'
    for title,body in [('Before (verified RC1)',before[name]),('After (diagnostic RC2)',after[name])]:
        excerpts+='\n'+title+'\n\n```csharp\n'+body+'\n```\n'
(OUT/'SOURCE_EXCERPTS.md').write_text(excerpts)
(OUT/'SOURCE_SHA256.txt').write_text(sha+'  Install/scripts/GlassHeroesBusiness.3.cs\n')
(OUT/'PR_BODY.md').write_text('''Diagnostic-only v0.3.4 RC2 recovery from owner-verified RC1 e312d8c9500f4a990eb993b82e8e76b55711d3de (same source at approval-only parent 75e3623f608a05d1a2b59d12c243dbcdfc061358). The October 24 0xC00000FD was observed during a Mechanic 1 customer-car transfer; responsible script/cause remain unknown.

Ten mechanic/customer transfer methods gain bounded entry/finally-exit breadcrumbs. Removing diagnostics restores the entire RC1 source exactly. No new native call/world scan, task/AI/route/financial/stock/clock change or save/settings key/default/migration. Separate 128 KiB trace + one previous segment, stable-state suppression, 12 ENTER admissions/ten seconds and existing safe-writer backoff.

Actual net48/C#7.3 Release against SHVDN3 3.6.0: 0 errors/two inherited warnings. Full 954 inherited + 31 diagnostic tests: 985/985 behavior/navigation PASS; static 21/21, inventory structure 13/13 and exact-restoration/fixture preservation PASS. GTA/Windows execution NOT RUN. No crash fix claimed.

See docs/features/v0.3.4-rc2-diagnostic for exact source diff, investigation limitations, Word/PDF developer guide, raw test evidence, diagnostic capture, GTA checks and rollback. Owner implementation scope pre-authorized; finished package awaits owner review. Keep draft; do not merge into main or the RC1 checkpoint/feature branch. No deployment or v0.3.5 work.
''')
documents=['DEVELOPER_GUIDE.md','DIAGNOSTIC_INSTRUCTIONS.md','INSTALL.md','GTA_TEST_CHECKLIST.md','ROLLBACK.md','CHANGELOG.md','SOURCE_EXCERPTS.md']
text='\n\n'.join((OUT/name).read_text() for name in documents)
for name,file in [('RC2Text','DejaVuSans.ttf'),('RC2Bold','DejaVuSans-Bold.ttf'),('RC2Mono','DejaVuSansMono.ttf')]:
    pdfmetrics.registerFont(TTFont(name,'/usr/share/fonts/truetype/dejavu/'+file))
styles=getSampleStyleSheet()
for name in ['Normal','BodyText','Title','Heading2']:
    styles[name].fontName='RC2Bold' if name in ['Title','Heading2'] else 'RC2Text'
styles.add(ParagraphStyle(name='RC2Code',fontName='RC2Mono',fontSize=7,leading=9))
doc=Document();doc.styles['Normal'].font.size=Pt(10)
flow=[];code=False
for line in text.splitlines():
    if line.startswith('```'):code=not code;continue
    if code:
        run=doc.add_paragraph().add_run(line);run.font.name='Courier New';run.font.size=Pt(8)
        for part in textwrap.wrap(line,94,replace_whitespace=False,drop_whitespace=False) or ['']:
            flow.append(Preformatted(part,styles['RC2Code']))
    elif not line:flow.append(Spacer(1,5))
    elif line.startswith('#'):
        level=len(line)-len(line.lstrip('#'));title=line[level:].strip()
        doc.add_heading(title,level=min(level-1,3));flow.append(Paragraph(escape(title),styles['Title' if level==1 else 'Heading2']))
    else:doc.add_paragraph(line);flow.append(Paragraph(escape(line),styles['BodyText']))
stem='GlassHeroesBusinessManager_v0.3.4_RC2_Diagnostic_Experimental_Developer_Guide'
doc.save(OUT/(stem+'.docx'))
def footer(canvas,document):
    canvas.setFont('RC2Text',8);canvas.drawString(42,22,'v0.3.4 RC2 | DIAGNOSTIC-ONLY | GTA NOT RUN | cause UNATTRIBUTED');canvas.drawRightString(570,22,str(document.page))
pdf=OUT/(stem+'.pdf')
SimpleDocTemplate(str(pdf),pagesize=(612,792),leftMargin=42,rightMargin=42,topMargin=40,bottomMargin=38).build(flow,onFirstPage=footer,onLaterPages=footer)
reader=PdfReader(pdf)
for content in ['\n'.join(p.text for p in Document(OUT/(stem+'.docx')).paragraphs),'\n'.join(p.extract_text() for p in reader.pages)]:
    for term in [sha,BASE,CODE,'985/985','954/954','Save.ini','Settings.ini','DIAGNOSTIC-ONLY','NOT RUN','UNATTRIBUTED']:
        assert term in content,term
receipt={'instrumentation_commit':CODE,'rc1_source_commit':BASE,'rc2_source_sha256':sha,'source_diff_reproduced':True,'word_pdf_text_checked':True,'pdf_pages':len(reader.pages),'gta_testing':'NOT RUN','crash_cause':'UNATTRIBUTED','crash_fix_claimed':False,'owner_review':'PENDING'}
(OUT/'REPORT_RECEIPT.json').write_text(json.dumps(receipt,indent=2)+'\n')
print(json.dumps(receipt,indent=2))
